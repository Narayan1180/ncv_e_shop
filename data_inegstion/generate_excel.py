import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from pathlib import Path
import random


# ============================================================
# CONFIG
# ============================================================

OUTPUT_PATH = (
    "/Users/ncvhome/NodeApp/Ecommerce/core/"
    "data_ingestion/data/sensor_hierarchy_6000_rows.xlsx"
)

NUM_ROWS = 6000


# ============================================================
# SENSOR PATHS
# ============================================================

sensor_paths = [
    "/se1/2021/water_sensor/t1/spectrum",
    "/se1/2021/water_sensor/t2/spectrum",
    "/se1/2021/water_sensor/t3/spectrum",

    "/se1/2021/fire_sensor/f1/spectrum",
    "/se1/2021/fire_sensor/f2/spectrum",
    "/se1/2021/fire_sensor/f3/spectrum",
]


# Add 1000 other sensors
for i in range(1, 1001):
    sensor_paths.append(
        f"/se1/2021/other_sensor_{i:04d}/spectrum"
    )


# ============================================================
# CREATE WORKBOOK
# ============================================================

output_path = Path(OUTPUT_PATH)
output_path.parent.mkdir(parents=True, exist_ok=True)

wb = Workbook()
ws = wb.active
ws.title = "sensor_data"


# ============================================================
# CREATE HEADERS
#
# Sensor 1:
#
# A1:B1 -> MERGED sensor path
# A2     -> f
# B2     -> v
# C      -> blank
#
# Sensor 2:
#
# D1:E1 -> MERGED sensor path
# D2     -> f
# E2     -> v
# F      -> blank
# ============================================================

for sensor_index, sensor_path in enumerate(sensor_paths):

    start_col = sensor_index * 3 + 1

    freq_col = start_col
    val_col = start_col + 1

    # --------------------------------------------------------
    # Merge sensor header across f + v
    # --------------------------------------------------------

    ws.merge_cells(
        start_row=1,
        start_column=freq_col,
        end_row=1,
        end_column=val_col
    )

    header_cell = ws.cell(
        row=1,
        column=freq_col,
        value=sensor_path
    )

    header_cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    header_cell.font = Font(bold=True)

    # --------------------------------------------------------
    # f / v
    # --------------------------------------------------------

    ws.cell(
        row=2,
        column=freq_col,
        value="f"
    )

    ws.cell(
        row=2,
        column=val_col,
        value="v"
    )

    ws.cell(
        row=2,
        column=freq_col
    ).alignment = Alignment(horizontal="center")

    ws.cell(
        row=2,
        column=val_col
    ).alignment = Alignment(horizontal="center")


# ============================================================
# GENERATE 6000 ROWS
# ============================================================

random.seed(42)

for i in range(NUM_ROWS):

    excel_row = i + 3

    # Frequency
    frequency = round(i * 0.1, 4)

    for sensor_index in range(len(sensor_paths)):

        start_col = sensor_index * 3 + 1

        # ----------------------------------------------------
        # f
        # ----------------------------------------------------

        ws.cell(
            row=excel_row,
            column=start_col,
            value=frequency
        )

        # ----------------------------------------------------
        # v
        # ----------------------------------------------------

        value = round(
            random.uniform(0, 100),
            4
        )

        ws.cell(
            row=excel_row,
            column=start_col + 1,
            value=value
        )

        # ----------------------------------------------------
        # start_col + 2 is intentionally blank
        # ----------------------------------------------------


# ============================================================
# COLUMN WIDTHS
# ============================================================

for sensor_index in range(len(sensor_paths)):

    start_col = sensor_index * 3 + 1

    ws.column_dimensions[
        openpyxl.utils.get_column_letter(start_col)
    ].width = 12

    ws.column_dimensions[
        openpyxl.utils.get_column_letter(start_col + 1)
    ].width = 12

    # Gap column
    ws.column_dimensions[
        openpyxl.utils.get_column_letter(start_col + 2)
    ].width = 3


# ============================================================
# ROW HEIGHT
# ============================================================

ws.row_dimensions[1].height = 30


# ============================================================
# SAVE
# ============================================================

wb.save(output_path)


# ============================================================
# OUTPUT
# ============================================================

print("========================================")
print("Excel generated successfully!")
print("========================================")
print(f"File       : {output_path}")
print(f"Sensors    : {len(sensor_paths)}")
print(f"Rows/sensor: {NUM_ROWS}")
print(f"Data rows  : {NUM_ROWS}")
print("========================================")